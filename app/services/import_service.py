import csv
import io
import json
from datetime import datetime, timezone

from app.extensions import db
from app.models.financial import BulkImportBatch, BulkImportRow, Transaction
from app.services.financial_service import compute_fingerprint, find_duplicates

REQUIRED_COLUMNS = ['type', 'amount', 'payee', 'date', 'location_id', 'payment_method']


def parse_csv(file_content, filename):
    reader = csv.DictReader(io.StringIO(file_content))
    rows = list(reader)
    return rows, reader.fieldnames or []


def parse_excel(file_obj):
    import openpyxl
    wb = openpyxl.load_workbook(file_obj, read_only=True)
    ws = wb.active
    headers = [str(cell.value or '').strip().lower() for cell in ws[1]]
    rows = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        row_dict = {}
        for i, val in enumerate(row):
            if i < len(headers):
                row_dict[headers[i]] = str(val) if val is not None else ''
        rows.append(row_dict)
    wb.close()
    return rows, headers


def preview_import(rows, columns, user_id, filename):
    batch = BulkImportBatch(
        filename=filename,
        uploaded_by=user_id,
        row_count=len(rows),
        status='previewing',
    )
    db.session.add(batch)
    db.session.flush()

    error_count = 0
    duplicate_count = 0

    for idx, row_data in enumerate(rows, start=1):
        errors = []
        for col in REQUIRED_COLUMNS:
            if not row_data.get(col, '').strip():
                errors.append(f'Missing required field: {col}')

        amount = row_data.get('amount', '')
        try:
            float(amount)
        except (ValueError, TypeError):
            errors.append(f'Invalid amount: {amount}')

        if row_data.get('type', '').lower() not in ('income', 'expense'):
            errors.append(f'Type must be income or expense')

        status = 'pending'
        error_msg = None
        matched_id = None

        if errors:
            status = 'error'
            error_msg = '; '.join(errors)
            error_count += 1
        else:
            fp = compute_fingerprint(
                row_data.get('payee', ''),
                row_data.get('amount', '0'),
                row_data.get('date', ''),
                row_data.get('location_id', ''),
                row_data.get('payment_method', ''),
            )
            dupes = find_duplicates(fp)
            if dupes:
                status = 'duplicate'
                matched_id = dupes[0][0].id
                duplicate_count += 1

        import_row = BulkImportRow(
            batch_id=batch.id,
            row_number=idx,
            raw_data=json.dumps(row_data),
            status=status,
            error_message=error_msg,
            matched_transaction_id=matched_id,
        )
        db.session.add(import_row)

    batch.error_count = error_count
    batch.duplicate_count = duplicate_count
    db.session.commit()
    return batch


def execute_import(batch_id, user_id):
    batch = db.session.get(BulkImportBatch, batch_id)
    if not batch or batch.status != 'previewing':
        return None, 'Batch not in preview state.'

    imported = 0
    for row in batch.rows.filter(BulkImportRow.status.in_(['pending', 'duplicate'])).all():
        if row.status == 'duplicate' and row.resolution == 'skip':
            continue
        if row.status == 'duplicate' and not row.resolution:
            continue

        data = row.get_data()
        txn = Transaction(
            type=data.get('type', 'income').lower(),
            amount=float(data.get('amount', 0)),
            payee=data.get('payee', ''),
            description=data.get('description', ''),
            category=data.get('category', ''),
            payment_method=data.get('payment_method', ''),
            location_id=int(data.get('location_id', 1)),
            handler_id=user_id,
            status='active',
        )
        txn.compute_fingerprint()
        db.session.add(txn)
        row.status = 'imported'
        imported += 1

    batch.status = 'imported'
    db.session.commit()
    return batch, f'{imported} transactions imported.'
